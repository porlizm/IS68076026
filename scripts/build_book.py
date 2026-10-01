# -*- coding: utf-8 -*-
"""
build_book.py — สร้างเล่ม .docx จาก book/*.md (ห้ามแก้ docx ด้วยมือ) · เขียนใหม่ 1 ต.ค. 2569 ตาม Prompt_Report หัวข้อ 9

  python scripts/book_numbers.py && python scripts/build_book.py      -> build/IS_68076026_Final_<DDMMMYY>.docx
  python scripts/build_book.py --check                                 -> ตรวจ placeholder · การอ้างอิง · ภาพ · สมการ (ไม่สร้างไฟล์)
  python scripts/export_pdf.py                                         -> PDF ผ่าน LibreOffice (อัปเดตสารบัญก่อนส่งออก)

ขั้นตอน
  1 แทน {{key}} จาก book/numbers.json · ภาพจาก book/figures/figures.json (ความกว้าง · alt text)
  2 แทน [@key] ด้วยเลข IEEE ตามลำดับที่อ้างครั้งแรก (บทที่ 1–3 แล้วภาคผนวก) · รายการอ้างอิงเรียงตามเลข · รายการที่ไม่มีใครอ้างไม่พิมพ์
  3 แปลงเป็น pandoc markdown ที่ใช้ custom-style · แทรก ZWSP ระหว่างคำไทยด้วย pythainlp (newmm)
  4 pandoc → docx (สมการ OMML · สารบัญ/สารบัญตาราง/สารบัญรูปเป็นฟิลด์) แล้วปรับด้วย python-docx
     A4 · ขอบบน/ล่าง/ขวา 2.54 ซม. ซ้าย 3.81 ซม. · TH Sarabun New ทุก run (ascii/hAnsi/cs) · sz=szCs · th-TH bidi
     บทที่ N กับชื่อบทคนละบรรทัด 20pt · H2 18pt · H3 16pt · เนื้อความ 16pt ย่อหน้า 1.65 ซม. กระจายแบบไทย
     ตาราง 14pt กว้างไม่เกิน 14.65 ซม. หัวตารางซ้ำทุกหน้า · ชื่อตารางเหนือตาราง · ชื่อรูปใต้รูปกึ่งกลาง
     สมการกึ่งกลางเลขสมการชิดขวา · เลขหน้ามุมขวาบน ส่วนหน้าใช้ ก ข ค · ฝังฟอนต์ 4 ไฟล์จาก assets/fonts
ต้องมี pandoc ≥ 2.9, python-docx, pythainlp
"""
import argparse, datetime, glob, io, json, os, re, shutil, subprocess, sys, tempfile, uuid, zipfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
BOOK = os.path.join(ROOT, "book")
OUT = os.path.join(ROOT, "build")
FONTS = os.path.join(ROOT, "assets", "fonts")
PARTS = ["01_chapter1.md", "02_chapter2.md", "03_chapter3.md", "04_references.md", "05_appendix.md"]
REFS = "04_references.md"
FONT = "TH Sarabun New"
TEXT_W_CM = 14.65
MAX_FIG_H = 13.5  # ซม.
KEEP_ROWS = 12    # ตารางที่สั้นกว่านี้ให้อยู่หน้าเดียวกันทั้งตาราง
TWIP = 566.929  # twips ต่อ 1 ซม.
PAGEBREAK = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```'
NBSP2 = "  "

from pythainlp.tokenize import word_tokenize


def zwsp(s):
    """แทรก ZWSP ระหว่างคำไทย (เฉพาะช่วงอักษรไทย) เพื่อให้ Word/LibreOffice ตัดบรรทัดตรงขอบคำ"""
    out = []
    for chunk in re.split(r"([฀-๿]+)", s):
        if chunk and re.match(r"[฀-๿]", chunk):
            out.append("​".join(w for w in word_tokenize(chunk, engine="newmm", keep_whitespace=True) if w))
        else:
            out.append(chunk)
    return "".join(out)


def div(style, text):
    return f'::: {{custom-style="{style}"}}\n{text}\n:::'


def field(instr, style="สารบัญ"):
    return ('```{=openxml}\n<w:p><w:pPr><w:pStyle w:val="' + style + '"/></w:pPr><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> ' + instr + ' </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>(คลิกขวาแล้วเลือกอัปเดตฟิลด์)</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>\n```')


def section_break(fmt, start=True):
    """ปิดส่วน (section) ปัจจุบัน · fmt = รูปแบบเลขหน้าของส่วนที่ปิด (None = ไม่มีเลขหน้า) · start = เริ่มนับ 1 ใหม่"""
    pg = (f'<w:pgNumType w:fmt="{fmt}"' + (' w:start="1"' if start else "") + "/>") if fmt else ""
    return ('```{=openxml}\n<w:p><w:pPr><w:spacing w:before="0" w:after="0"/><w:sectPr><w:type w:val="nextPage"/>'
            '<w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="2160" w:header="720" w:footer="720" w:gutter="0"/>'
            + pg + '</w:sectPr></w:pPr></w:p>\n```')


# ---------------------------------------------------------------- ข้อมูลเข้า
def load_json(*p):
    return json.load(open(os.path.join(ROOT, *p), encoding="utf-8"))


def resolve(text, nums, fname):
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)  # หมายเหตุในไฟล์ต้นฉบับไม่เข้าเล่ม
    def rep(m):
        k = m.group(1)
        if k not in nums: raise SystemExit(f"{fname}: ไม่มีค่า {{{{{k}}}}} ใน book/numbers.json (รัน python scripts/book_numbers.py)")
        return str(nums[k])
    return re.sub(r"\{\{([a-zA-Z0-9_]+)\}\}", rep, text)


def escape_md(t):
    t = re.sub(r"⏳\s*รอข้อมูล:\s*([^\n]+)", r"[รอข้อมูล: \1]", t)  # ⏳ ไม่มีในฟอนต์ไทย · คงเครื่องหมายรอข้อมูลไว้ในวงเล็บ
    return re.sub(r"(?<!\\)O\*NET", r"O\\*NET", t)


def citations(texts):
    """คืน (ลำดับ key, dict key->เลข) ตามการอ้างครั้งแรก"""
    order = []
    for f in PARTS:
        if f == REFS: continue
        for k in re.findall(r"\[@([A-Za-z0-9_]+)\]", texts[f]):
            if k not in order: order.append(k)
    return order, {k: i + 1 for i, k in enumerate(order)}


def references(text, num):
    entries = dict(re.findall(r"^\[@([A-Za-z0-9_]+)\]\s+(.+)$", text, re.M))
    missing = [k for k in num if k not in entries]
    if missing: raise SystemExit(f"อ้าง key ที่ไม่มีใน {REFS}: {missing}")
    unused = [k for k in entries if k not in num]
    out = [f"[{num[k]}] {entries[k]}" for k in sorted(num, key=num.get)]
    return out, unused


# ---------------------------------------------------------------- แปลงบล็อก
class Conv:
    def __init__(self, figs):
        self.figs = figs; self.alts = []; self.eq = 0; self.h1 = 0

    def block(self, b):
        b = b.strip()
        if not b or b.startswith("<!--"): return None
        if b.startswith("# "):
            # แต่ละบทเป็น section ของตัวเอง เพื่อไม่แสดงเลขหน้าในหน้าแรกของบท (เลขหน้าต่อเนื่อง)
            self.h1 += 1
            pb = "" if self.h1 == 1 else section_break("decimal", start=self.h1 == 2) + "\n\n"
            return pb + "# " + b[2:].strip()
        if b.startswith("## ภาคผนวก"): return PAGEBREAK + "\n\n" + b
        if b.startswith("#"): return b
        if b.startswith("$$"):
            m = re.match(r"\$\$\s*(.*?)\s*\\tag\{([^}]*)\}\s*\$\$$", b, re.S)
            if not m: raise SystemExit(f"สมการไม่มีเลข: {b[:60]}")
            self.eq += 1
            return div("สมการ", f"$${m.group(1)}$$") + "\n\n" + div("เลขสมการ", f"({m.group(2)})")
        m = re.match(r"\*\*(ตารางที่ [^*]+)\*\*\s*(.*)$", b, re.S)
        if m: return div("ชื่อตาราง", f"**{m.group(1)}**{NBSP2}{zwsp(m.group(2))}")
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)$", b)
        if m:
            key = os.path.splitext(os.path.basename(m.group(2)))[0]
            meta = self.figs.get(key)
            if not meta: raise SystemExit(f"ไม่มี {key} ใน book/figures/figures.json (รัน python scripts/make_figures.py)")
            self.alts.append(m.group(1) or meta["alt"])
            w = min(float(meta["width_cm"]), TEXT_W_CM)
            # จำกัดความสูงรูปไม่เกิน MAX_FIG_H ซม. เพื่อไม่ให้เหลือหน้าว่างครึ่งหน้า แต่ตัวอักษรเมื่อพิมพ์ต้องไม่ต่ำกว่า 12 pt
            px_w, px_h = meta.get("px", [1, 1])
            if w * px_h / px_w > MAX_FIG_H:
                w_h = MAX_FIG_H * px_w / px_h
                w_font = float(meta["width_cm"]) * 12.0 / float(meta.get("print_font_pt", 12) or 12)
                w = max(w_h, min(w, w_font))
            return div("รูปภาพ", f"![]({os.path.join(BOOK, m.group(2))}){{width={w:.2f}cm}}")
        m = re.match(r"\*(รูปที่ [^*]+)\*\s*(.*)$", b, re.S)
        if m: return div("ชื่อรูป", f"**{m.group(1)}**{NBSP2}{zwsp(m.group(2))}")
        if b.startswith("|"):
            return "\n".join(zwsp(r) for r in b.split("\n") if r.strip())
        if b.startswith("ที่มา "):
            return div("ที่มา", "**ที่มา** " + zwsp(b[len("ที่มา "):]))
        if b.startswith("[") and re.match(r"\[\d+\] ", b):
            return div("เอกสารอ้างอิง", b)
        b = re.sub(r"^(\d+)\. ", r"\1\\. ", b)
        return div("เนื้อหาทั่วไป", zwsp(b))


def front_matter(nums):
    t = resolve(open(os.path.join(BOOK, "00_front.md"), encoding="utf-8").read(), nums, "00_front.md")
    _, y, body = t.split("---", 2)
    meta = dict(re.findall(r'^(\w+):\s*"(.*)"\s*$', y, re.M))
    body = escape_md(re.sub(r"<!--.*?-->", "", body, flags=re.S))
    th = [("ปก", meta["title_th"]), ("ปกเว้น", ""), ("ปกเล็ก", meta["author_th"]), ("ปกเล็ก", "รหัสนักศึกษา " + meta["student_id"]),
          ("ปกเว้น", ""), ("ปกเล็ก", "อาจารย์ที่ปรึกษา " + meta["advisor_th"]), ("ปกเว้น", ""),
          ("ปกเล็ก", meta["course_th"]), ("ปกเล็ก", meta["program_th"]), ("ปกเล็ก", meta["major_th"]),
          ("ปกเล็ก", meta["faculty_th"]), ("ปกเล็ก", meta["institute_th"]), ("ปกเล็ก", meta["term_th"])]
    en = [("ปก", meta["title_en"].upper()), ("ปกเว้น", ""), ("ปกเล็ก", meta["author_en"].upper()), ("ปกเล็ก", "STUDENT ID " + meta["student_id"]),
          ("ปกเว้น", ""), ("ปกเล็ก", "ADVISOR " + meta["advisor_en"].upper()), ("ปกเว้น", ""),
          ("ปกเล็ก", meta["course_en"]), ("ปกเล็ก", meta["program_en"]), ("ปกเล็ก", meta["major_en"]),
          ("ปกเล็ก", meta["faculty_en"]), ("ปกเล็ก", meta["institute_en"]), ("ปกเล็ก", meta["term_en"])]
    cov = lambda rows: "\n\n".join(div(s, zwsp(x) if x else " ") for s, x in rows)
    parts = [cov(th), PAGEBREAK, cov(en), section_break(None)]
    secs = [s for s in re.split(r"^# ", body.strip(), flags=re.M) if s.strip()]
    toc_done = False
    for i, sec in enumerate(secs):
        head, _, rest = sec.partition("\n")
        head = head.strip()
        if head.startswith("คำอธิบายคำย่อ") and not toc_done:  # สารบัญอยู่ก่อนคำอธิบายคำย่อ
            parts += toc_pages(); toc_done = True
        blocks = [div("หัวเรื่องหน้าต้น", head)]
        for p in re.split(r"\n\s*\n", rest.strip()):
            p = p.strip()
            if not p: continue
            if p.startswith("|"): blocks.append("\n".join(zwsp(r) for r in p.split("\n") if r.strip()))
            elif p.startswith("<p align"): blocks.append(div("ลงชื่อ", re.sub(r"<[^>]+>", "", p)))
            elif re.match(r"\*\*(คำสำคัญ|Keywords)", p): blocks.append(div("คำสำคัญ", zwsp(p)))
            else: blocks.append(div("เนื้อหาทั่วไป", zwsp(p)))
        parts.append(("" if i == 0 else PAGEBREAK + "\n\n") + "\n\n".join(blocks))
    if not toc_done: parts += toc_pages()
    parts.append(section_break("thaiLetters"))
    return "\n\n".join(parts), meta


def toc_pages():
    return [PAGEBREAK + "\n\n" + div("หัวเรื่องหน้าต้น", "สารบัญ") + "\n\n" + field('TOC \\o "1-3" \\h \\z \\u'),
            PAGEBREAK + "\n\n" + div("หัวเรื่องหน้าต้น", "สารบัญตาราง") + "\n\n" + field('TOC \\h \\z \\t "ชื่อตาราง,1"'),
            PAGEBREAK + "\n\n" + div("หัวเรื่องหน้าต้น", "สารบัญรูป") + "\n\n" + field('TOC \\h \\z \\t "ชื่อรูป,1"')]


# ---------------------------------------------------------------- reference.docx
def reference_docx(path):
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    with open(path, "wb") as fh:
        subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], check=True, stdout=fh)
    d = docx.Document(path)
    sty = lambda n: next((x for x in d.styles if x.name == n), None)

    def pstyle(name, size, bold=False, align=None, indent=None, before=0, after=0, keep=False, char=False):
        st = sty(name) or d.styles.add_style(name, WD_STYLE_TYPE.CHARACTER if char else WD_STYLE_TYPE.PARAGRAPH)
        set_font(st.element, size, bold)
        if char: return st
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0
        if align is not None: pf.alignment = align
        pf.first_line_indent = indent if indent is not None else Cm(0)
        pf.left_indent = Cm(0)
        if keep: pf.keep_with_next = True
        return st

    for n in ["Normal", "Body Text", "First Paragraph", "Compact"]: pstyle(n, 16, align=A.LEFT)
    pstyle("Heading 1", 20, True, A.CENTER, after=18, keep=True)
    pstyle("Heading 2", 18, True, A.LEFT, before=12, after=6, keep=True)
    pstyle("Heading 3", 16, True, A.LEFT, before=6, after=3, keep=True)
    for h in ("Heading 1", "Heading 2", "Heading 3"): sty(h).font.color.rgb = None; sty(h).font.italic = False
    thai_justify(pstyle("เนื้อหาทั่วไป", 16, False, A.JUSTIFY, Cm(1.65)))
    pstyle("ชื่อตาราง", 16, False, A.LEFT, before=6, after=3, keep=True)
    pstyle("ชื่อรูป", 16, False, A.CENTER, before=3, after=9)
    pstyle("รูปภาพ", 16, False, A.CENTER, before=6, keep=True)
    pstyle("ที่มา", 14, False, A.LEFT, before=3, after=6)
    pstyle("สมการ", 16, False, A.CENTER, before=3, after=3)
    pstyle("เลขสมการ", 16, False, A.RIGHT)
    pstyle("หัวเรื่องหน้าต้น", 20, True, A.CENTER, after=18)
    pstyle("ปก", 20, True, A.CENTER, after=6)
    pstyle("ปกเล็ก", 18, True, A.CENTER)
    pstyle("ปกเว้น", 16, False, A.CENTER, after=48)
    pstyle("ลงชื่อ", 16, False, A.RIGHT, before=24)
    pstyle("คำสำคัญ", 16, False, A.LEFT, before=12)
    pstyle("สารบัญ", 16, False, A.LEFT)
    ref = pstyle("เอกสารอ้างอิง", 16, False, A.LEFT, Cm(-1.0), after=6)
    ref.paragraph_format.left_indent = Cm(1.0)
    for n in ("toc 1", "toc 2", "toc 3", "table of figures"):
        st = sty(n) or d.styles.add_style(n, WD_STYLE_TYPE.PARAGRAPH)
        set_font(st.element, 16, False)
    for st in d.styles:  # ตัวอักษรทุกสไตล์ (รวม Verbatim Char · Hyperlink) เป็น TH Sarabun New
        try: set_font(st.element, None, None)
        except Exception: pass
    from docx.oxml.ns import qn as _qn
    for rf in d.styles.element.iter(_qn("w:rFonts")):  # รวม docDefaults · ลบ theme font ทุกแบบ (asciiTheme, cstheme ฯลฯ)
        for k in list(rf.attrib):
            if "heme" in k: del rf.attrib[k]
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(_qn(k), FONT)
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3.81), Cm(2.54), Cm(2.54), Cm(2.54)
    s.header_distance = Cm(1.27)
    d.save(path)


def thai_justify(st):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    ppr = st.element.get_or_add_pPr()
    for jc in ppr.findall(qn("w:jc")): ppr.remove(jc)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "thaiDistribute"); ppr.append(jc)


def set_font(el, size, bold):
    """ตั้ง rFonts ascii/hAnsi/cs/eastAsia = TH Sarabun New · ลบ theme font · sz=szCs · lang bidi th-TH"""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    rpr = el.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr"); el.append(rpr)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in list(rf.attrib):
        if a.endswith("Theme"): del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    if size is not None:
        for tag in ("w:sz", "w:szCs"):
            e = rpr.find(qn(tag))
            if e is None: e = OxmlElement(tag); rpr.append(e)
            e.set(qn("w:val"), str(int(size * 2)))
    if bold is not None:
        for tag in ("w:b", "w:bCs"):
            e = rpr.find(qn(tag))
            if bold and e is None: rpr.append(OxmlElement(tag))
            if not bold and e is not None: rpr.remove(e)
    lang = rpr.find(qn("w:lang"))
    if lang is None: lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "en-US"); lang.set(qn("w:bidi"), "th-TH")


# ---------------------------------------------------------------- ปรับหลัง pandoc
def postprocess(path, alts, meta):
    import docx
    from docx.shared import Pt, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    d = docx.Document(path)
    body = d.element.body
    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

    # 1) บทที่ N ขึ้นบรรทัดใหม่ก่อนชื่อบท
    for p in d.paragraphs:
        if p.style.name == "Heading 1" and re.match(r"บทที่\s*\d+\s", p.text.replace("​", "")):
            r0 = p.runs[0]
            m = re.match(r"(บทที่\s*\d+)\s+(.*)$", "".join(r.text for r in p.runs).replace("​", ""))
            for r in p.runs[1:]: r._r.getparent().remove(r._r)
            r0.text = m.group(1); r0.add_break(); r0.add_text(m.group(2))

    # 2) สมการ: ตาราง 1×3 ไร้เส้น · สมการกึ่งกลาง · เลขชิดขวา
    for p in [x for x in d.paragraphs if x.style.name == "สมการ"]:
        el = p._p; nxt = el.getnext()
        if el.find(".//" + M + "oMath") is None or nxt is None: raise SystemExit("สมการแปลงไม่สำเร็จ")
        tbl = equation_table()
        el.addprevious(tbl)
        tcs = tbl.findall(".//" + qn("w:tc"))
        tcs[0].append(empty_p()); tcs[1].append(el); tcs[2].append(nxt)

    # 3) ตารางเนื้อหา: กว้าง 14.65 ซม. · คอลัมน์ตามความยาวข้อความ · เส้นขอบ · หัวตารางซ้ำ · 14pt
    for t in d.tables:
        cap = t._tbl.tblPr.find(qn("w:tblCaption"))
        if cap is not None and cap.get(qn("w:val")) == "equation": continue  # ตารางจัดสมการ
        format_table(t)

    # 4) alt text ของรูปตามลำดับ
    pics = body.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")
    if len(pics) != len(alts): print(f"คำเตือน: รูปใน docx {len(pics)} ไม่เท่ากับ alt {len(alts)}")
    for dp, alt in zip(pics, alts): dp.set("descr", alt); dp.set("title", alt[:60])

    # 5) ทุก run ใช้ TH Sarabun New (ยกเว้นสมการ) · sz=szCs
    for r in body.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None: rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
        rf = rpr.find(qn("w:rFonts"))
        if rf is None: rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in list(rf.attrib): del rf.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
        sz = rpr.find(qn("w:sz")); szcs = rpr.find(qn("w:szCs"))
        if sz is not None and szcs is None:
            szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), sz.get(qn("w:val"))); sz.addnext(szcs)
        b = rpr.find(qn("w:b"))
        if b is not None and rpr.find(qn("w:bCs")) is None: b.addnext(OxmlElement("w:bCs"))
    # rPr ต้องเรียงตาม schema: rStyle ก่อน rFonts
    for rpr in body.iter(qn("w:rPr")):
        rs = rpr.find(qn("w:rStyle"))
        if rs is not None and rpr.index(rs) != 0: rpr.remove(rs); rpr.insert(0, rs)

    # 6) ส่วน (section): ปก ไม่มีเลขหน้า · ส่วนหน้า ก ข ค · เนื้อหา 1 2 3 · เลขหน้ามุมขวาบน
    secs = d.sections
    for i, s in enumerate(secs):
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3.81), Cm(2.54), Cm(2.54), Cm(2.54)
        s.header_distance = Cm(1.27)
        s.header.is_linked_to_previous = False
        hp = s.header.paragraphs[0]
        for r in list(hp.runs): r._r.getparent().remove(r._r)
        if i == 0: continue
        if i >= 2:  # หน้าแรกของบท: ไม่แสดงเลขหน้า
            s.different_first_page_header_footer = True
            s.first_page_header.is_linked_to_previous = False
            for r in list(s.first_page_header.paragraphs[0].runs): r._r.getparent().remove(r._r)
        hp.alignment = A.RIGHT
        r = hp.add_run()
        for typ, txt in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
            if typ:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ); r._r.append(fc)
            elif txt == " PAGE ":
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
            else:
                tt = OxmlElement("w:t"); tt.text = txt; r._r.append(tt)
        r.font.name = FONT; r.font.size = Pt(16)
        set_font(r._r, 16, False)
    last = secs[-1]._sectPr
    pg = last.find(qn("w:pgNumType"))
    if pg is None: pg = OxmlElement("w:pgNumType"); last.append(pg)
    pg.set(qn("w:fmt"), "decimal")
    if len(secs) == 3: pg.set(qn("w:start"), "1")
    elif qn("w:start") in pg.attrib: del pg.attrib[qn("w:start")]
    # pgNumType ต้องอยู่หลัง pgMar ตาม schema
    pm = last.find(qn("w:pgMar"))
    if pm is not None: last.remove(pg); pm.addnext(pg)

    # 7) ค่าตั้งเอกสาร: อัปเดตฟิลด์เมื่อเปิด · ฝังฟอนต์ · ภาษาไทย
    st = d.settings.element
    for tag, val in (("w:embedTrueTypeFonts", None), ("w:updateFields", "true")):
        e = OxmlElement(tag)
        if val: e.set(qn("w:val"), val)
        st.insert(0, e) if tag == "w:embedTrueTypeFonts" else st.append(e)
    tfl = st.find(qn("w:themeFontLang"))
    if tfl is None: tfl = OxmlElement("w:themeFontLang"); st.append(tfl)
    tfl.set(qn("w:val"), "en-US"); tfl.set(qn("w:bidi"), "th-TH")
    d.core_properties.title = meta["title_th"]
    d.core_properties.author = meta["author_en"]
    d.core_properties.subject = "IS 68076026"
    d.core_properties.language = "th-TH"
    d.save(path)
    embed_fonts(path)


def empty_p():
    from docx.oxml import OxmlElement
    return OxmlElement("w:p")


def equation_table():
    from docx.oxml import parse_xml
    w = [1000, 6306, 1000]  # รวม 8306 twips = 14.65 ซม.
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    cells = "".join(f'<w:tc><w:tcPr><w:tcW w:w="{x}" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr></w:tc>' for x in w)
    grid = "".join(f'<w:gridCol w:w="{x}"/>' for x in w)
    return parse_xml(f'<w:tbl {ns}><w:tblPr><w:tblW w:w="{sum(w)}" w:type="dxa"/><w:jc w:val="center"/>'
                     '<w:tblBorders><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/>'
                     '<w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders><w:tblLayout w:type="fixed"/>'
                     '<w:tblLook w:val="0000"/><w:tblCaption w:val="equation"/></w:tblPr>'
                     f'<w:tblGrid>{grid}</w:tblGrid><w:tr>{cells}</w:tr></w:tbl>')


def text_cm(s):
    """ความกว้างโดยประมาณ (ซม.) ที่ 14pt TH Sarabun New"""
    s = re.sub(r"[\u200b\u0e31\u0e34-\u0e3a\u0e47-\u0e4e]", "", s)
    return sum(0.21 if "\u0e00" <= ch <= "\u0e7f" else (0.25 if ch.isupper() else 0.17) for ch in s)


def text_len(s):
    s = s.replace("​", "")
    return len(re.sub(r"[ัิ-ฺ็-๎]", "", s))  # ไม่นับสระบน/ล่างและวรรณยุกต์


def format_table(t):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    tbl = t._tbl; tblPr = tbl.tblPr
    total = int(round(TEXT_W_CM * TWIP))
    rows = tbl.findall(qn("w:tr"))
    grid_txt = [["".join(x.text or "" for x in tc.iter(qn("w:t"))) for tc in tr.findall(qn("w:tc"))] for tr in rows]
    ncol = max(len(r) for r in grid_txt)
    # ความกว้างคอลัมน์ (ซม.): ขั้นต่ำ = คำที่ยาวที่สุดไม่ต้องตัดกลางคำ · ส่วนที่เหลือแบ่งตามความยาวข้อความ
    PAD, CAP, SHORT = 0.4, 5.0, 2.2  # ระยะขอบเซลล์ · เพดานคำยาว (URL/แฮช) · คอลัมน์สั้นที่ไม่ย่อ
    mins, need = [], []
    for j in range(ncol):
        cells = [r[j] for r in grid_txt if j < len(r)]
        lens = [text_cm(c) for c in cells] or [0.2]
        tok = max((text_cm(w) for c in cells for w in re.split(r"[\s\u200b]+", c) if w), default=0.2)
        mins.append(min(tok + PAD, CAP))
        need.append(max(0.5 * max(lens) + 0.5 * sum(lens) / len(lens) + PAD, mins[-1]))
    W = TEXT_W_CM
    if sum(mins) >= W:
        fixed = [m <= SHORT for m in mins]
        rest = W - sum(m for m, f in zip(mins, fixed) if f)
        sm = sum(m for m, f in zip(mins, fixed) if not f) or 1; sn = sum(n for n, f in zip(need, fixed) if not f) or 1
        cm = [m if f else rest * (0.6 * m / sm + 0.4 * n / sn) for m, n, f in zip(mins, need, fixed)]
    else:
        rem = W - sum(mins); extra = [n - m for n, m in zip(need, mins)]
        if sum(extra) > rem:
            cm = [m + rem * e / sum(extra) for m, e in zip(mins, extra)]
        else:
            left = rem - sum(extra)
            cm = [n + left * n / sum(need) for n in need]
    widths = [int(x * TWIP) for x in cm]
    widths[-1] += total - sum(widths)
    for tag in ("w:tblW", "w:tblBorders", "w:tblLayout", "w:jc", "w:tblInd"):
        for e in tblPr.findall(qn(tag)): tblPr.remove(e)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(total)); tw.set(qn("w:type"), "dxa")
    sty = tblPr.find(qn("w:tblStyle"))
    (sty.addnext(tw) if sty is not None else tblPr.insert(0, tw))
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tw.addnext(jc)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); b.append(e)
    jc.addnext(b)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); b.addnext(lay)
    for e in tblPr.findall(qn("w:tblCellMar")): tblPr.remove(e)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "57"); e.set(qn("w:type"), "dxa"); mar.append(e)  # 0.1 ซม.
    lay.addnext(mar)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid"); tblPr.addnext(grid)
    for gc in list(grid): grid.remove(gc)
    for x in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(x)); grid.append(gc)
    small = len(rows) <= KEEP_ROWS and sum(len(c) for r in grid_txt for c in r) < 700  # ตารางสั้นไม่แยกหน้า
    for ri, tr in enumerate(rows):
        trPr = tr.find(qn("w:trPr"))
        if trPr is None:
            trPr = OxmlElement("w:trPr"); tr.insert(0 if tr.find(qn("w:tblPrEx")) is None else 1, trPr)
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0: trPr.append(OxmlElement("w:tblHeader"))
        for j, tc in enumerate(tr.findall(qn("w:tc"))):
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is None: tcPr = OxmlElement("w:tcPr"); tc.insert(0, tcPr)
            for e in tcPr.findall(qn("w:tcW")): tcPr.remove(e)
            cw = OxmlElement("w:tcW"); cw.set(qn("w:w"), str(widths[min(j, ncol - 1)])); cw.set(qn("w:type"), "dxa"); tcPr.insert(0, cw)
            for pe in tc.findall(qn("w:p")):
                p = Paragraph(pe, t._parent)
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
                if ri == 0: p.alignment = A.CENTER
                if (small and ri < len(rows) - 1) or ri == 0: p.paragraph_format.keep_with_next = True
                for r in p.runs:
                    r.font.size = Pt(14)
                    if ri == 0: r.font.bold = True


def embed_fonts(path):
    """ฝัง TH Sarabun New 4 ไฟล์แบบ obfuscated font ตาม ECMA-376 (fontTable + .odttf)"""
    files = {"embedRegular": "THSarabunNew.ttf", "embedBold": "THSarabunNew Bold.ttf",
             "embedItalic": "THSarabunNew Italic.ttf", "embedBoldItalic": "THSarabunNew BoldItalic.ttf"}
    missing = [f for f in files.values() if not os.path.exists(os.path.join(FONTS, f))]
    if missing: print("คำเตือน: ไม่พบฟอนต์สำหรับฝัง", missing); return
    zin = zipfile.ZipFile(path); items = {n: zin.read(n) for n in zin.namelist()}; zin.close()
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ft = items["word/fontTable.xml"].decode("utf-8")
    rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    embeds = []
    for i, (tag, fn) in enumerate(files.items(), 1):
        g = str(uuid.uuid4()).upper(); key = "{" + g + "}"
        kb = bytes.fromhex(g.replace("-", ""))[::-1]
        data = bytearray(open(os.path.join(FONTS, fn), "rb").read())
        for k in range(32): data[k] ^= kb[k % 16]
        items[f"word/fonts/font{i}.odttf"] = bytes(data)
        rels.append(f'<Relationship Id="rIdF{i}" Type="{R}/font" Target="fonts/font{i}.odttf"/>')
        embeds.append(f'<w:{tag} r:id="rIdF{i}" w:fontKey="{key}"/>')
    rels.append("</Relationships>")
    items["word/_rels/fontTable.xml.rels"] = "\n".join(rels).encode("utf-8")
    fontel = (f'<w:font w:name="{FONT}"><w:panose1 w:val="020B0500040200020003"/><w:charset w:val="00"/>'
              '<w:family w:val="swiss"/><w:pitch w:val="variable"/>' + "".join(embeds) + "</w:font>")
    ft = re.sub(r'<w:font w:name="' + re.escape(FONT) + r'">.*?</w:font>', "", ft, flags=re.S)
    if "xmlns:r=" not in ft.split(">", 2)[1]:
        ft = re.sub(r"<w:fonts\b", f'<w:fonts xmlns:r="{R}"', ft, count=1)
    ft = ft.replace("</w:fonts>", fontel + "</w:fonts>")
    items["word/fontTable.xml"] = ft.encode("utf-8")
    ct = items["[Content_Types].xml"].decode("utf-8")
    if 'Extension="odttf"' not in ct:
        ct = ct.replace("<Types ", "<Types ", 1).replace(
            "</Types>", '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/></Types>')
    items["[Content_Types].xml"] = ct.encode("utf-8")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n in ["[Content_Types].xml"] + [n for n in items if n != "[Content_Types].xml"]: z.writestr(n, items[n])
    os.replace(tmp, path)


# ---------------------------------------------------------------- main
def build_markdown():
    nums = load_json("book", "numbers.json")
    figs = load_json("book", "figures", "figures.json")
    texts = {f: escape_md(resolve(open(os.path.join(BOOK, f), encoding="utf-8").read(), nums, f)) for f in PARTS}
    order, num = citations(texts)
    refs, unused = references(texts[REFS], num)
    conv = Conv(figs)
    front, meta = front_matter(nums)
    body = []
    for f in PARTS:
        t = texts[f]
        if f == REFS:
            body.append(conv.block("# เอกสารอ้างอิง"))
            body += [conv.block(r) for r in refs]
            continue
        t = re.sub(r"\[@([A-Za-z0-9_]+)\]", lambda m: f"[{num[m.group(1)]}]", t)
        for blk in re.split(r"\n\s*\n", t):
            c = conv.block(blk)
            if c: body.append(c)
    md = front + "\n\n" + "\n\n".join(body) + "\n"
    info = dict(citations=len(order), unused_refs=unused, figures=len(conv.alts), equations=conv.eq,
                tables=md.count('custom-style="ชื่อตาราง"'), figure_captions=md.count('custom-style="ชื่อรูป"'))
    left = re.findall(r"\{\{[^}]*\}\}|\[@[^\]]*\]", md)
    if left: raise SystemExit(f"ยังมี placeholder ค้าง: {left[:5]}")
    imgs = re.findall(r"!\[\]\(([^)]+)\)", md)
    miss = [i for i in imgs if not os.path.exists(i)]
    if miss: raise SystemExit(f"ภาพหาย: {miss}")
    return md, meta, conv.alts, info


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--check", action="store_true"); ap.add_argument("--date"); a = ap.parse_args()
    md, meta, alts, info = build_markdown()
    print("ตรวจ:", " · ".join(f"{k} {v}" for k, v in info.items()))
    if info["figures"] != info["figure_captions"]: raise SystemExit("จำนวนรูปกับชื่อรูปไม่เท่ากัน")
    if a.check: return
    os.makedirs(OUT, exist_ok=True)
    tag = a.date or datetime.date.today().strftime("%d%b%y").upper()
    out = os.path.join(OUT, f"IS_68076026_Final_{tag}.docx")
    with tempfile.TemporaryDirectory() as tmp:
        ref = os.path.join(tmp, "reference.docx"); reference_docx(ref)
        src = os.path.join(tmp, "book.md"); open(src, "w", encoding="utf-8").write(md)
        subprocess.run(["pandoc", src, "-f", "markdown+pipe_tables+tex_math_dollars+fenced_divs+raw_attribute-implicit_figures",
                        "-t", "docx", "--reference-doc", ref, "-o", out], check=True)
    postprocess(out, alts, meta)
    shutil.copy(out, os.path.join(OUT, "IS_68076026_Final_latest.docx"))
    json.dump(dict(info, file=os.path.relpath(out, ROOT)), open(os.path.join(OUT, "build_info.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("เขียน", os.path.relpath(out, ROOT))


if __name__ == "__main__":
    main()
